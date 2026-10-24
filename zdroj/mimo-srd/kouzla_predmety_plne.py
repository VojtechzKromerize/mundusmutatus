#!/usr/bin/env python3
"""Plné znění kouzel a předmětů (Vojtovy tabulky Kouzla.xlsx a Předměty.xlsx) → JSON pro tabulku pravidla_texty
(klíče „kouzla“ a „predmety“), zobrazuje se jen přihlášeným (mm_pravidla_text / mm_pravidla_seznam).

Texty NEJSOU v repozitáři (je veřejný) ani v art/ – zdrojem jsou Vojtovy tabulky. Převod:
  - sjednotí pojmy (pojmy_textu.py: AC, vypravěč, pozice kouzla 3. stupně, vyčerpání, školy, vlastnosti zbraní, typy zranění),
  - připojí rámečky Artikulí (kouzla_predmety_artikule.py) jako odstavce „> **Artikule – Název.** text“,
  - vyřadí položky, které do roku 1458 nepatří, a moderní zbraně ze seznamů základních předmětů,
  - české názvy vezme z voleb webu (zdroj/kouzla/jmena.json, zdroj/vybaveni/jmena.json), jinak z tabulky.

  python3 zdroj/mimo-srd/kouzla_predmety_plne.py Kouzla.xlsx Předměty.xlsx VÝSTUPNÍ_SLOŽKA [--xlsx]
    → VÝSTUP/kouzla.json, VÝSTUP/predmety.json (objekt id → položka), s --xlsx i upravené tabulky pro Vojtu.
Nahrání do databáze: viz docs/STAV.md („Plné texty kouzel a předmětů“) – šifrovaný soubor + pg_net + pgp_sym_decrypt.
"""
import json
import os
import re
import sys

D = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, D)
from pojmy_textu import pojmy, mistrovstvi  # noqa: E402
import kouzla_predmety_artikule as A  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(D))
SRDJM = {"evard's black tentacles": "Black Tentacles", "melf's acid arrow": "Acid Arrow", "mordenkainen's sword": "Arcane Sword",
         "bigby's hand": "Arcane Hand", "tasha's hideous laughter": "Hideous Laughter", "otto's irresistible dance": "Irresistible Dance",
         "leomund's tiny hut": "Tiny Hut", "otiluke's resilient sphere": "Resilient Sphere", "mordenkainen's private sanctum": "Private Sanctum",
         "nystul's magic aura": "Arcanist's Magic Aura", "rary's telepathic bond": "Telepathic Bond", "drawmij's instant summons": "Instant Summons",
         "mordenkainen's faithful hound": "Faithful Hound", "leomund's secret chest": "Secret Chest", "tenser's floating disk": "Floating Disk",
         "otiluke's freezing sphere": "Freezing Sphere", "mordenkainen's magnificent mansion": "Magnificent Mansion"}


def jm(n):
    """„English (Česky)“ → (en, cz); český název je poslední závorka (smí obsahovat vnořené závorky)"""
    n = str(n).strip()
    if n.endswith(")"):
        hl = 0
        for i in range(len(n) - 1, -1, -1):
            if n[i] == ")":
                hl += 1
            elif n[i] == "(":
                hl -= 1
                if hl == 0:
                    return n[:i].strip(), n[i + 1:-1].strip()
    return n, ""


def radky(cesta):
    import openpyxl
    ws = openpyxl.load_workbook(cesta, read_only=True).worksheets[0]
    R = list(ws.iter_rows(values_only=True))
    return [dict(zip(R[0], r)) for r in R[1:] if any(c is not None for c in r)]


def prirazeni(en):
    """rámečky Artikulí pro položku: (název, text)"""
    out, kody = [], []
    for k, L in A.PRIRAZENI.items():
        if k == en or (k.endswith("*") and en.startswith(k[:-1])):
            kody += [c for c in L if c not in kody]
    for c in kody:
        out.append(A.ZASADY[c])
    for n, t in A.VLASTNI.get(en, []):
        out = [o for o in out if o[0] != n] + [(n, t)]
    return out


def ramecky(L):
    return "".join(f"\n\n> **Artikule – {n}.** {t}" for n, t in L)


def zakladni(t):
    """seznamy základních předmětů: bez moderních zbraní, Pistol = píšťala, Musket = hákovnice"""
    out = []
    for r in t.split("\n"):
        if re.search(A.MODERNI, r):
            continue
        r = re.sub(r"^(\s*(?:•\s*)?)Pistole\b", r"\1Píšťala", r)
        r = re.sub(r"^(\s*(?:•\s*)?)Mušketa\b", r"\1Hákovnice", r)
        out.append(r)
    return "\n".join(out)


def kouzla(cesta):
    web = {w["en"].lower(): w for w in json.load(open(os.path.join(ROOT, "art", "kouzla.json"), encoding="utf-8"))}
    jmena = json.load(open(os.path.join(ROOT, "zdroj", "kouzla", "jmena.json"), encoding="utf-8"))
    out, pouzite = {}, set()
    for x in radky(cesta):
        en, cz = jm(x["Název"])
        if en in A.VYRADIT:
            continue
        wen = en if en.lower() in web else SRDJM.get(en.lower(), en)
        w = web.get(wen.lower())
        sk = str(x["Škola"] or "")
        ct = str(x["Doba sesílání"] or "")
        if w and ct in ("Reakce", "Bonusová akce") and w["ct"].startswith(ct) and len(w["ct"]) > len(ct):
            ct = w["ct"]  # spouštěč reakce či bonusové akce (tabulka ho nemá)
        if en == "Plant Growth":
            ct = "Akce (Bujný porost) nebo 8 hodin (Obohacení)"
        ur = str(x["Úroveň"] or "")
        sl = re.sub(r"\bR \(", "M (", str(x["Složky"] or ""))
        t = pojmy(x["Popis"] or "") + ramecky(prirazeni(en) or prirazeni(wen))
        k = wen.lower()
        if k in out:
            raise SystemExit(f"CHYBA: kouzlo {en} je v tabulce dvakrát")
        out[k] = {"en": wen, "cz": jmena.get(wen) or pojmy(cz), "l": 0 if ur.startswith("Trik") else int(ur[0]),
                  "s": sk.replace(" (rituál)", ""), "ri": 1 if "(rituál)" in sk else 0, "ct": pojmy(ct), "r": pojmy(x["Dosah"] or ""),
                  "k": pojmy(sl), "d": pojmy(x["Trvání"] or ""), "t": t.strip(), "v": pojmy(x["Při seslání z vyšší pozice / růst triku"] or "")}
        if wen != en:
            out[k]["en2"] = en
        pouzite.add(k)
    return out


def predmety(cesta):
    jmena = {k.lower(): v for k, v in json.load(open(os.path.join(ROOT, "zdroj", "vybaveni", "jmena.json"), encoding="utf-8")).items()}
    out = {}
    for x in radky(cesta):
        en, cz = jm(x["Název"])
        if en in A.VYRADIT or str(x["Typ"] or "") in A.VYRADIT_TYP:
            continue
        t = pojmy(zakladni(x["Popis"] or ""))
        if en.startswith("Spell Gem (Amber)"):
            t = t.replace("+9", "+7")
        if en.startswith("Spell Gem (Topaz)"):
            t = t.replace("+10", "+9")
        o = {"en": en, "cz": jmena.get(en.lower()) or pojmy(cz), "vz": x["Vzácnost"] or "", "typ": pojmy(x["Typ"] or ""),
             "sl": pojmy(x["Připoutání"] or ""), "zr": pojmy(x["Zranění"] or ""), "vl": pojmy(x["Vlastnosti"] or ""),
             "ma": mistrovstvi(x["Weapon Mastery"] or ""), "w": x["Hmotnost"] or "", "t": (t + ramecky(prirazeni(en))).strip()}
        o = {k: v for k, v in o.items() if v not in ("", None)}
        k = en.lower()
        if k in out:
            if out[k].get("t") == o.get("t"):
                continue  # stejná položka dvakrát (Prosthetic Limb)
            n = 2
            while f"{k} ({n})" in out:
                n += 1
            k = f"{k} ({n})"
        out[k] = o
    return out


def xlsx(src, dst, data, druh):
    """upravená tabulka pro Vojtu: sjednocené pojmy, sloupec Artikule, bez vyřazených položek"""
    import openpyxl
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    hl = [c.value for c in ws[1]]
    ws.cell(1, len(hl) + 1, "Artikule")
    smaz = []
    for i in range(2, ws.max_row + 1):
        n = ws.cell(i, 1).value
        if not n:
            continue
        en, _ = jm(n)
        if druh == "predmety" and (en in A.VYRADIT or str(ws.cell(i, hl.index("Typ") + 1).value or "") in A.VYRADIT_TYP) or druh == "kouzla" and en in A.VYRADIT:
            smaz.append(i)
            continue
        for j, h in enumerate(hl, 1):
            v = ws.cell(i, j).value
            if isinstance(v, str) and h != "Název":
                ws.cell(i, j).value = mistrovstvi(v) if h == "Weapon Mastery" else pojmy(zakladni(v) if h == "Popis" else v)
        R = prirazeni(en) or prirazeni(SRDJM.get(en.lower(), en))
        ws.cell(i, len(hl) + 1, "\n".join(f"{a}: {b}" for a, b in R) or None)
    for i in reversed(smaz):
        ws.delete_rows(i)
    wb.save(dst)


if __name__ == "__main__":
    a = [x for x in sys.argv[1:] if not x.startswith("--")]
    if len(a) != 3:
        raise SystemExit(__doc__)
    os.makedirs(a[2], exist_ok=True)
    K, P = kouzla(a[0]), predmety(a[1])
    for n, d in (("kouzla", K), ("predmety", P)):
        json.dump(d, open(os.path.join(a[2], n + ".json"), "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    nk = sum(1 for x in K.values() if "> **Artikule" in x.get("t", ""))
    np_ = sum(1 for x in P.values() if "> **Artikule" in x.get("t", ""))
    print(f"kouzla {len(K)} (s rámečkem Artikule {nk}), předměty {len(P)} (s rámečkem Artikule {np_})")
    if "--xlsx" in sys.argv:
        xlsx(a[0], os.path.join(a[2], "Kouzla.xlsx"), K, "kouzla")
        xlsx(a[1], os.path.join(a[2], "Předměty.xlsx"), P, "predmety")
        print("upravené tabulky:", os.path.join(a[2], "Kouzla.xlsx"), os.path.join(a[2], "Předměty.xlsx"))
